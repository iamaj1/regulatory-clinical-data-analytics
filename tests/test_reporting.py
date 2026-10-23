import io

import pandas as pd
from openpyxl import load_workbook

from src.reporting import csv_report, excel_report, pdf_report


def test_exports_are_generated():
    findings = pd.DataFrame([{
        "record_reference": "A",
        "rule_id": "X",
        "severity": "High",
        "explanation": "Test",
        "corrective_action": "Fix",
        "row": 1,
        "review_status": "Completed",
        "action_taken": "Corrected source data",
        "reviewer_notes": "Checked and corrected",
        "last_updated": "2026-10-23T12:00:00",
    }])
    assert b"rule_id" in csv_report(findings)
    assert len(excel_report(pd.DataFrame({"a": [1]}), findings)) > 100
    pdf = pdf_report("Report", findings)
    assert pdf.startswith(b"%PDF")
    assert len(pdf) > 500


def test_excel_includes_clinical_case_reviews_sheet():
    source = pd.DataFrame({"study_id": ["S1"], "subject_id": ["P001"]})
    findings = pd.DataFrame([{"rule_id": "CLN-002", "review_status": "Pending"}])
    cases = pd.DataFrame([{
        "Study ID": "S1",
        "Site ID": "SITE-1",
        "Subject ID": "P001",
        "Review status": "Completed",
        "Action taken": "Case review",
        "Reviewer notes": "Reviewed synthetic record",
    }])
    workbook = load_workbook(io.BytesIO(excel_report(source, findings, cases)), read_only=True)
    assert workbook.sheetnames == ["Source data", "Findings", "Case reviews"]
    assert workbook["Case reviews"].max_row == 2


def test_pdf_includes_case_review_content():
    findings = pd.DataFrame([{
        "rule_id": "CLN-002",
        "severity": "High",
        "record_reference": "S1 | SITE-1 | P001",
        "explanation": "Subject ID is missing.",
        "review_status": "In Review",
        "action_taken": "Requested clarification",
        "reviewer_notes": "Asked site to verify",
    }])
    cases = pd.DataFrame([{
        "Study ID": "S1",
        "Site ID": "SITE-1",
        "Subject ID": "P001",
        "Review status": "Completed",
        "Action taken": "Case review",
        "Reviewer notes": "Reviewed synthetic record",
    }])
    pdf = pdf_report("Clinical Review", findings, cases)
    assert pdf.startswith(b"%PDF")
    assert len(pdf) > 500
